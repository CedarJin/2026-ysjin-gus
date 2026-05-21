#!/usr/bin/env python3
"""Build the combined GUS protein database from the paper PDF and FMN table."""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path


PAGE_MARKER = re.compile(r"^--\s+\d+\s+of\s+\d+\s+--$")
HEADER = re.compile(r"^>(?P<protein_id>[^-\s]+)-(?P<clade>[A-Za-z0-9]+)\s*(?P<source>.*)$")
AA_LINE = re.compile(r"^[A-Za-zxX]+$")


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise SystemExit(
            "pypdf is required for PDF extraction. Install it in the analysis "
            "environment, or pass a text file converted from the PDF."
        ) from exc

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def parse_fasta_records(path: Path) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    seq_parts: list[str] = []

    def flush() -> None:
        nonlocal current, seq_parts
        if current is None:
            return
        sequence = "".join(seq_parts).upper().replace("X", "X")
        if not sequence:
            raise ValueError(f"Empty sequence for {current['reference_id']}")
        current["sequence"] = sequence
        current["length"] = str(len(sequence))
        records.append(current)
        current = None
        seq_parts = []

    with path.open() as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line or PAGE_MARKER.match(line):
                continue
            if line.startswith(">"):
                flush()
                match = HEADER.match(line)
                if not match:
                    raise ValueError(f"Could not parse FASTA header: {line}")
                protein_id = match.group("protein_id")
                clade = match.group("clade")
                source = " ".join(match.group("source").split())
                current = {
                    "reference_id": f"{protein_id}-{clade}",
                    "protein_id": protein_id,
                    "clade": clade,
                    "source": source,
                }
                continue
            if current is None:
                continue
            cleaned = re.sub(r"[^A-Za-zxX]", "", line)
            if cleaned and AA_LINE.match(cleaned):
                seq_parts.append(cleaned)

    flush()
    return records


def parse_text_records(text: str) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    seq_parts: list[str] = []

    def flush() -> None:
        nonlocal current, seq_parts
        if current is None:
            return
        sequence = "".join(seq_parts).upper().replace("X", "X")
        if not sequence:
            raise ValueError(f"Empty sequence for {current['reference_id']}")
        current["sequence"] = sequence
        current["length"] = str(len(sequence))
        records.append(current)
        current = None
        seq_parts = []

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line or PAGE_MARKER.match(line):
            continue
        if line.startswith(">"):
            flush()
            match = HEADER.match(line)
            if not match:
                raise ValueError(f"Could not parse FASTA header: {line}")
            protein_id = match.group("protein_id")
            clade = match.group("clade")
            source = " ".join(match.group("source").split())
            current = {
                "reference_id": f"{protein_id}-{clade}",
                "protein_id": protein_id,
                "clade": clade,
                "source": source,
            }
            continue
        if current is None:
            continue
        cleaned = re.sub(r"[^A-Za-zxX]", "", line)
        if cleaned and AA_LINE.match(cleaned):
            seq_parts.append(cleaned)

    flush()
    return records


def parse_fmn_records(path: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl is required to read Supplemental_Table_2.xlsx") from exc

    wb = load_workbook(path, data_only=True)
    ws = wb["GUS_Summary"]

    headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
    required = ["Gene_ID", "Gene_name", "Loop_Class", "Gene_Sequence"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise SystemExit(f"Missing expected columns in {path}: {', '.join(missing)}")

    records: list[dict[str, str]] = []
    for row in range(2, ws.max_row + 1):
        gene_id = ws.cell(row, headers["Gene_ID"]).value
        gene_name = ws.cell(row, headers["Gene_name"]).value
        loop_class = ws.cell(row, headers["Loop_Class"]).value
        sequence = ws.cell(row, headers["Gene_Sequence"]).value
        if not gene_id or not sequence:
            continue

        cleaned = re.sub(r"[^A-Za-zxX]", "", str(sequence)).upper()
        if not cleaned:
            continue
        records.append(
            {
                "reference_id": f"{gene_id}-FMN",
                "protein_id": str(gene_id),
                "clade": "FMN",
                "source": str(gene_name) if gene_name is not None else "",
                "loop_class": str(loop_class) if loop_class is not None else "",
                "sequence": cleaned,
                "length": str(len(cleaned)),
            }
        )

    return records


def write_outputs(records: list[dict[str, str]], fasta_path: Path, metadata_path: Path) -> None:
    fasta_path.parent.mkdir(parents=True, exist_ok=True)
    metadata_path.parent.mkdir(parents=True, exist_ok=True)

    with fasta_path.open("w") as fasta:
        for record in records:
            fasta.write(
                f">{record['reference_id']} clade={record['clade']} "
                f"source={record['source']}\n"
            )
            seq = record["sequence"]
            for i in range(0, len(seq), 80):
                fasta.write(seq[i : i + 80] + "\n")

    fields = ["reference_id", "protein_id", "clade", "source", "length"]
    with metadata_path.open("w") as metadata:
        metadata.write("\t".join(fields) + "\n")
        for record in records:
            metadata.write("\t".join(record[field] for field in fields) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--paper-pdf", type=Path)
    parser.add_argument("--paper-fasta", type=Path)
    parser.add_argument("--fmn-xlsx", required=True, type=Path)
    parser.add_argument("--fasta", required=True, type=Path)
    parser.add_argument("--metadata", required=True, type=Path)
    parser.add_argument("--expected-paper-count", type=int, default=279)
    args = parser.parse_args()

    if args.paper_fasta is not None:
        records = parse_fasta_records(args.paper_fasta)
    else:
        if args.paper_pdf is None:
            raise SystemExit("Provide --paper-pdf or --paper-fasta")
        records = parse_text_records(read_pdf(args.paper_pdf))
    if len(records) != args.expected_paper_count:
        raise SystemExit(
            f"Extracted {len(records)} paper records, expected {args.expected_paper_count}."
        )
    records.extend(parse_fmn_records(args.fmn_xlsx))
    write_outputs(records, args.fasta, args.metadata)
    print(f"Extracted {len(records)} records to {args.fasta}")


if __name__ == "__main__":
    main()
